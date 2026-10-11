"""Le carte intestate di prova (test/carta_intestata.js e la prova nel browser).

    python genera_carta.py      # scrive carta.docx e logo.png (serve python-docx e Pillow)

carta.pdf si fa col browser: node genera_carta_pdf.js (Playwright), dallo stesso disegno.
"""
import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw

QUI = os.path.dirname(os.path.abspath(__file__))
logo = Image.new('RGB', (400, 120), (13, 148, 136))
d = ImageDraw.Draw(logo)
d.rectangle([20, 20, 100, 100], fill=(255, 255, 255))
d.text((120, 45), 'STUDIO GEOLOGICO ROSSI', fill=(255, 255, 255))
logo.save(os.path.join(QUI, 'logo.png'))

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Mm(40), Mm(30), Mm(20), Mm(20)
s.header_distance, s.footer_distance = Mm(10), Mm(8)
h = s.header.paragraphs[0]
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
h.add_run().add_picture(os.path.join(QUI, 'logo.png'), width=Mm(80))
p = s.header.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Dott. Geol. Mario Rossi — Via Roma 1, Lecce')
r.bold = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
t = s.footer.add_table(rows=1, cols=3, width=Mm(170))
for c, testo in zip(t.rows[0].cells, ['Tel. 0832 000000', 'info@studiorossi.it', 'P.IVA 01234567890']):
    c.paragraphs[0].add_run(testo).font.size = Pt(8)
doc.add_paragraph('Corpo del documento (non fa parte della carta).')
doc.save(os.path.join(QUI, 'carta.docx'))
print('ok')
